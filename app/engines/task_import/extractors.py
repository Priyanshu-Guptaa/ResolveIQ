"""Format-specific row extraction for Structured Task Import.

Each extractor's only job is turning one source format's bytes into a
list of :class:`TaskRecord`. Nothing here touches persistence,
duplicate detection, relationships, or indexing -- that's all in
``importer.py``, shared identically regardless of which extractor ran.

Deliberately does NOT reuse ``app.engines.ingestion.parsers.xlsx_parser
.XlsxParser`` for the XLSX path: that parser flattens a whole sheet
into one pipe-delimited text blob for whole-document search (Knowledge
Management's upload flow) -- structured import needs column-name-keyed
access to each row individually, a genuinely different operation, not
a duplicate of that one.

Both extractors only emit rows that actually have resolution content
(the same "only import tasks with a captured resolution" filter the
user approved for the original JSON import) -- a row with no resolution
has nothing to teach the Recommendation Engine and would only dilute
search relevance.
"""

from __future__ import annotations

import io
import json
import logging
from pathlib import Path

from app.domain.task_import import TaskRecord

logger = logging.getLogger(__name__)

_STATE_LABELS = {"3": "Closed Complete", "4": "Closed Incomplete", "7": "Closed Skipped"}


def extract_tasks(filename: str, content: bytes) -> list[TaskRecord]:
    """Dispatches to the right extractor by file extension. Raises
    ValueError for anything else -- callers (the API endpoint) turn
    that into a 400, not a silent empty import."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        return extract_from_json_bytes(content)
    if suffix in (".xlsx", ".xlsm"):
        return extract_from_xlsx_bytes(content)
    raise ValueError(f"Unsupported task import format: {suffix or 'no extension'} (expected .json or .xlsx)")


def extract_from_json_bytes(content: bytes) -> list[TaskRecord]:
    """ServiceNow JSON export shape: ``{"records": [...]}`` (or a bare
    list, tolerated for other export tools). Field names are the raw
    ServiceNow internal names."""
    data = json.loads(content)
    raw_records = data.get("records", []) if isinstance(data, dict) else (data if isinstance(data, list) else [])

    records: list[TaskRecord] = []
    for item in raw_records:
        resolution = (item.get("close_notes") or "").strip()
        number = (item.get("number") or "").strip()
        if not resolution or not number:
            continue

        tags: list[str] = []
        priority = str(item.get("priority") or "")
        state = str(item.get("state") or "")
        if item.get("assignment_group"):
            tags.append(f"team:{item['assignment_group']}")
        if item.get("cmdb_ci"):
            tags.append(f"cmdb:{item['cmdb_ci']}")

        title = (item.get("short_description") or "").strip() or number
        records.append(
            TaskRecord(
                ticket_number=number,
                title=title[:500],
                description=(item.get("description") or "").strip() or "(no description captured)",
                resolution=resolution,
                priority=priority,
                state=_STATE_LABELS.get(state, state),
                extra_tags=tags,
            )
        )
    return records


# XLSX export's own column headers -- ground truth, not renamed to match
# the JSON export's internal field names, since the two formats are
# different reports pulled from different queries and mapping them onto
# identical labels would be a fabricated correspondence.
_XLSX_NUMBER_COL = "Number"
_XLSX_TITLE_COL = "Short description"
_XLSX_DESCRIPTION_COL = "Description"
_XLSX_RESOLUTION_COLS = ("Comments and Work notes", "Work notes")
"""Tried in order -- observed identical content in both columns on the
real export; falling back to the second covers a report variant where
only one of the two is populated."""
_XLSX_PRIORITY_COL = "Priority"
_XLSX_STATE_COL = "State"
_XLSX_ASSIGNEE_COL = "Assigned to"
_XLSX_LOCATION_COL = "Location"
_XLSX_TASK_TYPE_COL = "Task type"


def extract_from_xlsx_bytes(content: bytes) -> list[TaskRecord]:
    from openpyxl import load_workbook

    # read_only=False deliberately -- see xlsx_parser.py's docstring:
    # openpyxl's read_only streaming mode trusts a workbook's declared
    # <dimension> XML element, which real-world exports (this exact
    # ServiceNow report export included) can leave stale, silently
    # truncating iteration to a tiny fraction of the real data.
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    if not workbook.worksheets:
        return []
    sheet = workbook.worksheets[0]

    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        return []
    header = [str(cell).strip() if cell is not None else "" for cell in header_row]
    col_index = {name: i for i, name in enumerate(header)}

    def cell(row: tuple, column_name: str) -> str:
        idx = col_index.get(column_name)
        if idx is None or idx >= len(row):
            return ""
        value = row[idx]
        return "" if value is None else str(value).strip()

    records: list[TaskRecord] = []
    for row in rows:
        resolution = ""
        for col in _XLSX_RESOLUTION_COLS:
            resolution = cell(row, col)
            if resolution:
                break
        number = cell(row, _XLSX_NUMBER_COL)
        if not resolution or not number:
            continue

        tags: list[str] = []
        if assignee := cell(row, _XLSX_ASSIGNEE_COL):
            tags.append(f"assignee:{assignee}")
        if location := cell(row, _XLSX_LOCATION_COL):
            tags.append(f"location:{location}")
        if task_type := cell(row, _XLSX_TASK_TYPE_COL):
            tags.append(f"type:{task_type}")

        title = cell(row, _XLSX_TITLE_COL) or number
        records.append(
            TaskRecord(
                ticket_number=number,
                title=title[:500],
                description=cell(row, _XLSX_DESCRIPTION_COL) or "(no description captured)",
                resolution=resolution,
                priority=cell(row, _XLSX_PRIORITY_COL),
                state=cell(row, _XLSX_STATE_COL),
                extra_tags=tags,
            )
        )
    return records
