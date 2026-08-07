"""PowerPoint parser -- python-pptx (Sprint 3, Phase 3.2: Knowledge
Management's PPTX support). Same shape as DocxParser: a .pptx is a zip of
XML parts, so it must never be UTF-8 decoded directly.
"""

from __future__ import annotations

import io
import logging
from pathlib import Path

from app.domain.ingestion import FileKind, ParsedFile, ParseWarning

logger = logging.getLogger(__name__)


class PptxParser:
    def can_parse(self, filename: str) -> bool:
        return Path(filename).suffix.lower() == ".pptx"

    def parse(self, filename: str, content: bytes) -> ParsedFile:
        warnings: list[ParseWarning] = []
        text = ""
        metadata: dict = {}

        try:
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(content))
            slide_parts: list[str] = []
            notes_count = 0
            for slide_index, slide in enumerate(presentation.slides, start=1):
                slide_lines: list[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        slide_lines.append(shape.text_frame.text)
                    if shape.has_table:
                        for row in shape.table.rows:
                            slide_lines.append(" | ".join(cell.text for cell in row.cells))
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    slide_lines.append(f"[Speaker notes] {slide.notes_slide.notes_text_frame.text}")
                    notes_count += 1
                if slide_lines:
                    slide_parts.append(f"--- Slide {slide_index} ---\n" + "\n".join(slide_lines))

            text = "\n\n".join(slide_parts)
            metadata = {"slide_count": len(presentation.slides), "notes_slide_count": notes_count}
            if not text.strip():
                warnings.append(ParseWarning(message="Presentation contains no extractable text", severity="warning"))
        except Exception as exc:  # noqa: BLE001 -- a corrupt/unsupported .pptx must not crash the upload
            logger.warning("Failed to parse .pptx %s: %s", filename, exc)
            warnings.append(ParseWarning(message=f"Failed to parse .pptx: {exc}", severity="error"))

        return ParsedFile(filename=filename, kind=FileKind.PPTX, text=text, warnings=warnings, metadata=metadata)
