"""Spreadsheet parser -- openpyxl.

.xlsx is also a zip of XML parts (same reason .docx needed a real
parser). Cells are rendered as pipe-delimited rows per sheet, which keeps
tabular structure legible to both a human reading raw_content and the
regex-based entity extractor.
"""

from __future__ import annotations

import io
import logging
from pathlib import Path

from app.domain.ingestion import FileKind, ParsedFile, ParseWarning

logger = logging.getLogger(__name__)


class XlsxParser:
    def can_parse(self, filename: str) -> bool:
        return Path(filename).suffix.lower() in (".xlsx", ".xlsm")

    def parse(self, filename: str, content: bytes) -> ParsedFile:
        warnings: list[ParseWarning] = []
        text = ""
        metadata: dict = {}

        try:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            sections: list[str] = []
            for sheet in workbook.worksheets:
                sections.append(f"--- Sheet: {sheet.title} ---")
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(cell) for cell in row if cell is not None]
                    if cells:
                        sections.append(" | ".join(cells))
            text = "\n".join(sections)
            metadata = {"sheet_names": workbook.sheetnames}
            if not text.strip():
                warnings.append(ParseWarning(message="Workbook contains no non-empty cells", severity="warning"))
        except Exception as exc:  # noqa: BLE001
            logger.warning("Failed to parse .xlsx %s: %s", filename, exc)
            warnings.append(ParseWarning(message=f"Failed to parse .xlsx: {exc}", severity="error"))

        return ParsedFile(filename=filename, kind=FileKind.XLSX, text=text, warnings=warnings, metadata=metadata)
