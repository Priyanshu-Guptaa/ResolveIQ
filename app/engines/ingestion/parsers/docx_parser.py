"""Word document parser -- python-docx.

This is the direct fix for Problem 1's ".docx uploads display PK
headers": .docx is a zip of XML parts, so decoding it as UTF-8 text
produces exactly the "PK\\x03\\x04..." garbage observed in Sprint 1.
"""

from __future__ import annotations

import io
import logging
from pathlib import Path

from app.domain.ingestion import FileKind, ParsedFile, ParseWarning

logger = logging.getLogger(__name__)


class DocxParser:
    def can_parse(self, filename: str) -> bool:
        return Path(filename).suffix.lower() == ".docx"

    def parse(self, filename: str, content: bytes) -> ParsedFile:
        warnings: list[ParseWarning] = []
        text = ""
        metadata: dict = {}

        try:
            from docx import Document

            document = Document(io.BytesIO(content))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            table_rows = 0
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
                    table_rows += 1
            text = "\n".join(parts)
            metadata = {"paragraph_count": len(document.paragraphs), "table_row_count": table_rows}
            if not text.strip():
                warnings.append(ParseWarning(message="Document contains no extractable text", severity="warning"))
        except Exception as exc:  # noqa: BLE001 -- a corrupt/unsupported .docx must not crash the upload
            logger.warning("Failed to parse .docx %s: %s", filename, exc)
            warnings.append(ParseWarning(message=f"Failed to parse .docx: {exc}", severity="error"))

        return ParsedFile(filename=filename, kind=FileKind.DOCX, text=text, warnings=warnings, metadata=metadata)
