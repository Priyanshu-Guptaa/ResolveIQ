"""Tests for the Knowledge Management Engine (Sprint 3, Phase 3.2) --
the first Administration module.

Runs against real temp-file SQLite (same reasoning as
test_knowledge_foundation_migration.py: this is persistence/integration
behavior) and the *real* Evidence Ingestion Pipeline (IngestionEngine),
per the phase's explicit "use the existing pipeline, don't duplicate
parsing" requirement -- these tests are exactly what proves that reuse
actually works, not just that it was written. ChromaDB itself is faked
(FakeKnowledgeStore) so these stay fast and offline; indexing behavior
is verified by asserting what the fake was called with.
"""

from __future__ import annotations

import io
import tempfile
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from app.config import get_settings
from app.domain.enums import DocumentStatus, KnowledgeCollection
from app.domain.recommendation import KnowledgeMatch
from app.engines.ingestion.engine import IngestionEngine
from app.engines.ingestion.file_type_registry import FileTypeRegistry
from app.engines.knowledge.engine import KnowledgeEngine
from app.engines.knowledge_management.engine import DocumentNotFoundError, KnowledgeManagementEngine
from app.infrastructure.db.component_repository import SqlAlchemyComponentProfileRepository
from app.infrastructure.db.knowledge_repository import SqlAlchemyKnowledgeRepository
from app.infrastructure.db.seed_migration import migrate_component_profiles
from app.infrastructure.db.session import get_engine, get_session_factory

SAMPLE_DIR = get_settings().sample_knowledge_dir


class FakeKnowledgeStore:
    def __init__(self) -> None:
        self.upserts: dict[tuple, dict] = {}

    def upsert(self, collection, record_id, text, title, metadata) -> None:
        self.upserts[(collection, record_id)] = {"text": text, "title": title, "metadata": metadata}

    def query(self, collection, text, top_k: int = 5) -> list[KnowledgeMatch]:  # pragma: no cover
        return []

    def count(self, collection) -> int:
        return len([k for k in self.upserts if k[0] == collection])

    def list_recent(self, collection, limit: int = 5) -> list[KnowledgeMatch]:  # pragma: no cover
        return []

    def delete(self, collection, record_id: str) -> None:
        self.upserts.pop((collection, record_id), None)


@pytest.fixture
def setup():
    with tempfile.TemporaryDirectory() as tmp:
        db_path = Path(tmp) / "test.db"
        sqlite_url = f"sqlite:///{db_path.as_posix()}"
        session_factory = get_session_factory(sqlite_url)

        component_repo = SqlAlchemyComponentProfileRepository(session_factory)
        knowledge_repo = SqlAlchemyKnowledgeRepository(session_factory)
        migrate_component_profiles(component_repo, SAMPLE_DIR)  # seeds CommandProcessorHost etc.

        fake_store = FakeKnowledgeStore()
        knowledge_engine = KnowledgeEngine(fake_store, knowledge_repo)
        ingestion_engine = IngestionEngine(FileTypeRegistry())
        upload_dir = Path(tmp) / "uploads"

        engine = KnowledgeManagementEngine(
            knowledge_repo, component_repo, ingestion_engine, knowledge_engine, upload_dir
        )

        yield engine, knowledge_repo, fake_store, upload_dir

        get_engine(sqlite_url).dispose()
        get_engine.cache_clear()


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _pptx_bytes(title: str, body: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


# --- Upload + Extract ----------------------------------------------------


def test_upload_text_file_creates_draft_document(setup):
    engine, _, _, upload_dir = setup
    results = engine.upload_document("notes.txt", b"Plain text content.", uploaded_by="admin@resolveiq")

    assert len(results) == 1
    doc = results[0].document
    assert doc.status == DocumentStatus.DRAFT
    assert doc.content == "Plain text content."
    assert doc.title == "notes"
    assert doc.source == "upload"
    assert doc.created_by == "admin@resolveiq"
    assert doc.file_type == "text"
    assert doc.file_path is not None
    assert (upload_dir / doc.file_path).exists()
    assert (upload_dir / doc.file_path).read_bytes() == b"Plain text content."


def test_upload_docx_uses_evidence_ingestion_pipeline_not_a_duplicate_parser(setup):
    """Same DocxParser investigation evidence uploads already use --
    this is the direct proof of "reuse the pipeline, don't duplicate it.\""""
    engine, _, _, _ = setup
    content = _docx_bytes(["First paragraph.", "Mentions CommandProcessorHost explicitly."])
    results = engine.upload_document("playbook.docx", content)

    assert len(results) == 1
    doc = results[0].document
    assert doc.file_type == "docx"
    assert "CommandProcessorHost" in doc.content


def test_upload_pptx_extracts_slide_text(setup):
    engine, _, _, _ = setup
    content = _pptx_bytes("Meter Comm Triage", "Check collector logs first")
    results = engine.upload_document("triage.pptx", content)

    assert len(results) == 1
    doc = results[0].document
    assert doc.file_type == "pptx"
    assert "Meter Comm Triage" in doc.content
    assert "Check collector logs first" in doc.content


def test_upload_zip_fans_out_into_multiple_documents_sharing_one_original_file(setup):
    engine, _, _, upload_dir = setup
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("one.txt", "Content one")
        archive.writestr("two.txt", "Content two")

    results = engine.upload_document("bundle.zip", buf.getvalue())

    assert len(results) == 2
    titles = {r.document.title for r in results}
    assert titles == {"one", "two"}
    # Both documents point at the same saved original (the zip itself).
    file_paths = {r.document.file_path for r in results}
    assert len(file_paths) == 1
    assert (upload_dir / file_paths.pop()).exists()


def test_upload_unsupported_file_still_creates_a_document_with_a_warning(setup):
    engine, _, _, _ = setup
    results = engine.upload_document("firmware.bin", b"\x00\x01\x02\x03binarygarbage")

    assert len(results) == 1
    doc, warnings = results[0].document, results[0].warnings
    assert doc.status == DocumentStatus.DRAFT
    assert doc.content == ""  # never garbage-decoded, per the Ingestion Pipeline's core guarantee
    assert any(w.severity == "warning" for w in warnings)


# --- Metadata step -------------------------------------------------------


def test_update_metadata_sets_fields_and_leaves_others_unchanged(setup):
    engine, _, _, _ = setup
    doc = engine.upload_document("notes.txt", b"content")[0].document

    updated = engine.update_metadata(
        doc.id,
        title="GLP/DCW Triage Notes",
        tags=["glp", "dcw"],
        product="Command Center",
        technology="RF Mesh",
        related_components=["CommandProcessorHost"],
        updated_by="admin@resolveiq",
    )

    assert updated.title == "GLP/DCW Triage Notes"
    assert updated.tags == ["glp", "dcw"]
    assert updated.product == "Command Center"
    assert updated.technology == "RF Mesh"
    assert updated.related_components == ["CommandProcessorHost"]
    assert updated.version is None  # never touched, stays unset
    assert updated.status == DocumentStatus.DRAFT  # metadata edits don't change lifecycle
    assert updated.updated_by == "admin@resolveiq"


def test_update_metadata_on_published_document_reindexes(setup):
    engine, _, fake_store, _ = setup
    doc = engine.upload_document("notes.txt", b"original content")[0].document
    engine.publish_document(doc.id)

    engine.update_metadata(doc.id, title="Renamed Title")

    indexed = fake_store.upserts[(KnowledgeCollection.DOCUMENTATION, doc.id)]
    assert indexed["title"] == "Renamed Title"


def test_update_metadata_unknown_document_raises(setup):
    engine, _, _, _ = setup
    with pytest.raises(DocumentNotFoundError):
        engine.update_metadata("does-not-exist", title="x")


# --- Validate step ---------------------------------------------------------


def test_validate_flags_empty_content_missing_tags_and_missing_classification(setup):
    engine, _, _, _ = setup
    doc = engine.upload_document("firmware.bin", b"\x00\x01binary")[0].document

    warnings = engine.validate_document(doc.id)

    assert any("no extracted text" in w.lower() for w in warnings)
    assert any("no tags" in w.lower() for w in warnings)
    assert any("product, technology, or component" in w.lower() for w in warnings)


def test_validate_flags_duplicate_title(setup):
    engine, _, _, _ = setup
    first = engine.upload_document("guide.txt", b"first version")[0].document
    engine.update_metadata(first.id, title="Meter Comm Guide")
    second = engine.upload_document("guide2.txt", b"second version")[0].document
    engine.update_metadata(second.id, title="Meter Comm Guide")

    warnings = engine.validate_document(second.id)
    assert any("possible duplicate" in w.lower() for w in warnings)


def test_validate_clean_document_has_no_warnings(setup):
    engine, _, _, _ = setup
    doc = engine.upload_document("guide.txt", b"Useful triage content.")[0].document
    engine.update_metadata(doc.id, title="Unique Guide Title", tags=["triage"], product="Command Center")

    assert engine.validate_document(doc.id) == []


# --- Lifecycle ---------------------------------------------------------


def test_publish_indexes_into_the_knowledge_store(setup):
    engine, _, fake_store, _ = setup
    doc = engine.upload_document("guide.txt", b"Searchable triage content.")[0].document

    published = engine.publish_document(doc.id)

    assert published.status == DocumentStatus.PUBLISHED
    key = (KnowledgeCollection.DOCUMENTATION, doc.id)
    assert key in fake_store.upserts
    assert "Searchable triage content." in fake_store.upserts[key]["text"]


def test_archive_removes_a_published_document_from_the_index(setup):
    engine, _, fake_store, _ = setup
    doc = engine.upload_document("guide.txt", b"content")[0].document
    engine.publish_document(doc.id)
    assert (KnowledgeCollection.DOCUMENTATION, doc.id) in fake_store.upserts

    archived = engine.archive_document(doc.id)

    assert archived.status == DocumentStatus.ARCHIVED
    assert (KnowledgeCollection.DOCUMENTATION, doc.id) not in fake_store.upserts


def test_archiving_a_draft_document_never_touched_the_index_in_the_first_place(setup):
    engine, _, fake_store, _ = setup
    doc = engine.upload_document("guide.txt", b"content")[0].document

    archived = engine.archive_document(doc.id)

    assert archived.status == DocumentStatus.ARCHIVED
    assert (KnowledgeCollection.DOCUMENTATION, doc.id) not in fake_store.upserts


def test_restore_moves_an_archived_document_back_to_draft(setup):
    engine, _, fake_store, _ = setup
    doc = engine.upload_document("guide.txt", b"content")[0].document
    engine.publish_document(doc.id)
    engine.archive_document(doc.id)

    restored = engine.restore_to_draft(doc.id)

    assert restored.status == DocumentStatus.DRAFT
    assert (KnowledgeCollection.DOCUMENTATION, doc.id) not in fake_store.upserts


def test_get_document_returns_none_for_unknown_id(setup):
    engine, _, _, _ = setup
    assert engine.get_document("does-not-exist") is None


def test_lifecycle_actions_on_unknown_document_raise(setup):
    engine, _, _, _ = setup
    with pytest.raises(DocumentNotFoundError):
        engine.publish_document("does-not-exist")
    with pytest.raises(DocumentNotFoundError):
        engine.archive_document("does-not-exist")
    with pytest.raises(DocumentNotFoundError):
        engine.validate_document("does-not-exist")


# --- Dashboard -------------------------------------------------------------


def test_dashboard_stats_count_by_status(setup):
    engine, _, _, _ = setup
    draft = engine.upload_document("a.txt", b"a")[0].document
    to_publish = engine.upload_document("b.txt", b"b")[0].document
    to_archive = engine.upload_document("c.txt", b"c")[0].document
    engine.publish_document(to_publish.id)
    engine.publish_document(to_archive.id)
    engine.archive_document(to_archive.id)

    stats = engine.get_dashboard_stats()

    assert stats.total_count == 3
    assert stats.draft_count == 1
    assert stats.published_count == 1
    assert stats.archived_count == 1
    assert stats.under_review_count == 0
    assert {d.id for d in stats.recently_added} <= {draft.id, to_publish.id, to_archive.id}


# --- Knowledge Library: search / filter / sort / pagination ---------------


def test_list_documents_never_includes_content(setup):
    engine, _, _, _ = setup
    engine.upload_document("a.txt", b"some content that should never come back in a list view")

    page = engine.list_documents(page_size=10)
    assert len(page.items) == 1
    assert not hasattr(page.items[0], "content")  # DocumentationListItem has no content field at all


def test_list_documents_filters_by_status_and_product(setup):
    engine, _, _, _ = setup
    published = engine.upload_document("a.txt", b"a")[0].document
    engine.update_metadata(published.id, product="Command Center")
    engine.publish_document(published.id)
    engine.upload_document("b.txt", b"b")  # stays Draft

    published_only = engine.list_documents(status=DocumentStatus.PUBLISHED.value)
    assert [d.id for d in published_only.items] == [published.id]

    by_product = engine.list_documents(product="Command Center")
    assert [d.id for d in by_product.items] == [published.id]


def test_list_documents_search_matches_title(setup):
    engine, _, _, _ = setup
    doc = engine.upload_document("a.txt", b"a")[0].document
    engine.update_metadata(doc.id, title="GLP/DCW Mismatch Triage")
    engine.upload_document("b.txt", b"b")  # title stays "b"

    results = engine.list_documents(search="glp")
    assert [d.id for d in results.items] == [doc.id]


def test_list_documents_component_filter(setup):
    engine, _, _, _ = setup
    linked = engine.upload_document("a.txt", b"a")[0].document
    engine.update_metadata(linked.id, related_components=["CommandProcessorHost"])
    engine.upload_document("b.txt", b"b")  # not linked to any component

    results = engine.list_documents(component="CommandProcessorHost")
    assert [d.id for d in results.items] == [linked.id]
    assert results.items[0].related_components == ["CommandProcessorHost"]


def test_list_documents_pagination(setup):
    engine, _, _, _ = setup
    for i in range(5):
        engine.upload_document(f"doc{i}.txt", f"content {i}".encode())

    page1 = engine.list_documents(sort_by="title", sort_desc=False, page=1, page_size=2)
    page2 = engine.list_documents(sort_by="title", sort_desc=False, page=2, page_size=2)
    page3 = engine.list_documents(sort_by="title", sort_desc=False, page=3, page_size=2)

    assert page1.total_count == 5
    assert page1.total_pages == 3
    assert len(page1.items) == 2
    assert len(page2.items) == 2
    assert len(page3.items) == 1
    all_ids = {d.id for d in page1.items} | {d.id for d in page2.items} | {d.id for d in page3.items}
    assert len(all_ids) == 5  # no overlap, no gaps
